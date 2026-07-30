# Build the CharmEd Minds competency scoring-grid gradebook.
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

NAVY = "1F3864"; MINT = "1E7A63"; LIGHT = "EAF0F8"; YELLOW = "FFF6D5"
GREEN = "C6EFCE"; GREENT = "006100"; RED = "FFC7CE"; REDT = "9C0006"; GRAY = "F2F2F2"
AR = "Arial"

thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfill = PatternFill("solid", fgColor=NAVY)
hfont = Font(name=AR, bold=True, color="FFFFFF", size=10)
lfont = Font(name=AR, size=10)
bfont = Font(name=AR, bold=True, size=10)
yfill = PatternFill("solid", fgColor=YELLOW)
zfill = PatternFill("solid", fgColor=LIGHT)
wrap = Alignment(wrap_text=True, vertical="top")
ctr = Alignment(horizontal="center", vertical="center")

# ---- item bank: (item, part, question, correct answer, module, max) ----
items = [
 ("1","A","Central promise (one sentence)","“Confidence Through Cognition and Education”","1.2",1),
 ("2","A","Age range served","5–90","1.3",1),
 ("3","A","Name 3 of the 7 population tracks","Any 3: Children; Teens/YA; Adults; Neurodivergent Adults; Recovery; Chronic Disease; Seniors","1.3",1),
 ("4","A","T/F — only for diagnosing ADHD in children","False","1.1",1),
 ("5","A","Tier adding BHI + parent/family coaching","Gold","2.1",1),
 ("6","A","Three clinical pillars","Digital testing (CogniFit/Creyos); Quick Cognitive Test Kits; Weekly EF interventions (97127)","2.3",1),
 ("7","A","Two primary computerized batteries","CogniFit and Creyos","6A",1),
 ("8","A","Two guiding principles","Optimize the body before psychiatry; Behavior is communication","1.5",1),
 ("9","A","Order the 7 pathway steps","Intake → Screening → Medical review → Pattern ID → Education → Care plan → Follow-up","2.5",1),
 ("10","A","Most important step & why","Education — people need language for their experience before they can change it","2.5",1),
 ("11","A","Standard program length","12 weeks","2.4",1),
 ("12","B","Patient asks “do I have dementia?” — response & next step","Don’t diagnose; explain screening tools; Provider reviews & discusses; put on review list","3.6/9.2",1),
 ("13","B","T/F — coordinator may tell patient scores mean they have ADHD","False","3.6",1),
 ("14","B","Who may order & interpret labs","Provider","7.2",1),
 ("15","B","Two things staff MAY + two they MAY NOT do","MAY: administer screens, report scores, describe trends, educate, escalate. MAY NOT: diagnose, interpret clinically, adjust meds/order labs, decide medical necessity","3.6",1),
 ("16","B","Red flag (sudden decline) — what to do","Tell the APD immediately; APD routes to Provider the same day","8.4",1),
 ("17","B","Patient mentions self-harm — what to do","Follow the BHW crisis protocol immediately","8.4/10.5",1),
 ("18","B","Who assigns the diagnosis / ICD-10","Provider / NP","3.1",1),
 ("19","C","Max time a referral may be unassigned","24 hours","4.2",1),
 ("20","C","Home digital testing completed within","48 hours","4.2",1),
 ("21","C","Referral task to coordinator sent","Same day","4.2",1),
 ("22","C","Intake scheduled within","7 days","4.2",1),
 ("23","C","97127 session documented when","Same day","8.2",1),
 ("24","C","Match step → owner (4)","Baseline review=Provider; Assign coordinator=APD; Handout/schedule intake=MA; Runs intake=Care Coordinator","4.2",1),
 ("25","C","Monthly review — whose summary first & by whom","Coordinator’s summary, reviewed by the APD before the Provider","10.1",1),
 ("26","C","Quarterly full retest week","Week 12","10.1",1),
 ("27","C","3 of 5 continuation decisions","continue / modify tier / maintenance / discharge / refer","10.4",1),
 ("28","D","Match 5 tools → domain","ASRS=adult ADHD; STOP-Bang=OSA risk; RAADS-R=adult autism; PHQ-9=depression; AD8=caregiver decline","6",1),
 ("29","D","AuDHD clarification = which 4 categories","ADHD tool + autism-trait tool + executive-function tool + sensory profile","6D",1),
 ("30","D","T/F — a screen can replace a full autism dx evaluation","False","6D",1),
 ("31","D","Children’s ADHD screen + two informants","Vanderbilt; parent + teacher","6B",1),
 ("32","D","Quick cognitive screen for older adults","Mini-Cog","6I",1),
 ("33","D","Results combine into what document","The CharmEd Minds Cognitive Profile","6L",1),
 ("34","E","Five required elements of a 97127 note","ICD-10; deficits; interventions; functional response; homework","8.2",1),
 ("35","E","Code for provider interpretation","96132","8.1",1),
 ("36","E","Code for weekly training session","97127","8.1",1),
 ("37","E","Who decides medical necessity","Provider","8.1",1),
 ("38","E","APD note-audit frequency","Weekly","8.3",1),
 ("39","E","State the APD Ownership Principle","Faithful restatement: APD finds & resolves untracked patients, overdue notes, missed sessions, unassigned referrals — without being told","3.2",1),
 ("F1","F","Role-specific item 1 (for employee’s role)","Per role — see Answer Key F1–F5","3.x",1),
 ("F2","F","Role-specific item 2","Per role — see Answer Key F1–F5","3.x",1),
 ("F3","F","Role-specific item 3","Per role — see Answer Key F1–F5","3.x",1),
 ("53","G","Child snores/mouth-breathes — 2 sleep screens + who interprets","BEARS + Pediatric Sleep Questionnaire; Provider interprets the workup","6F/7",1),
 ("54","G","Reframe “he’s just lazy”","Brain-based EF / task-initiation framing; behavior is communication, not laziness","9.1",1),
 ("55","G","“Noncompliant” diabetic — interpretation + 1 support","Likely cognitively overloaded / medically unstable; support: simplify care plan (or med routine, monitoring, screen contributors)","Track 6",1),
 ("56","G","34yo masking/burnout — track + 1 battery category","Track 4 (late-identified neurodivergence); e.g., autism-trait (AQ/RAADS-R), CAT-Q, sensory profile, ASRS, BRIEF-A","Track 4",1),
 ("57","G","82yo memory change — 2 reversible contributors","Any 2: B12, thyroid, anemia, sleep apnea, depression, medication/anticholinergic, dehydration, infection, hearing/vision, BP","7.1",1),
]
parts = ["A","B","C","D","E","F","G"]
part_names = {"A":"Core Program Knowledge","B":"Scope of Practice & Safety (must be 100%)","C":"Patient Journey & Deadlines",
 "D":"Screening & the Cognitive Profile","E":"Documentation, Billing & Compliance","F":"Role-Specific Competency","G":"Application & Scenarios"}

wb = openpyxl.Workbook()

# ============ TAB 1: Instructions ============
ws = wb.active; ws.title = "Instructions"
ws.sheet_view.showGridLines = False
ws["B2"] = "CharmEd Minds™ — Competency Scoring Grid"; ws["B2"].font = Font(name=AR, bold=True, size=16, color=NAVY)
ws["B3"] = "Confidence Through Cognition and Education · BHW Medical Group"; ws["B3"].font = Font(name=AR, italic=True, size=10, color="555555")
lines = [
 ("How to use this workbook", True),
 ("1.  Score Sheet — grade one employee. Enter their details, then mark each item 1 (correct) or 0 (incorrect) in the yellow Awarded column. Totals, %, and the Result update automatically.", False),
 ("2.  Class Roster — track the whole team. Enter each person’s per-part points; the roster computes total, %, Part B check, and Result.", False),
 ("3.  Duplicate the Score Sheet tab (right-click → Move or Copy → check ‘Create a copy’) to grade additional employees.", False),
 ("", False),
 ("Passing standard", True),
 ("•  80% or higher overall, AND", False),
 ("•  100% on Part B (Scope of Practice & Safety) — every safety item correct.", False),
 ("•  If Part B is not perfect, the employee reviews Modules 3 & 8 and re-tests before working independently.", False),
 ("", False),
 ("Scoring notes", True),
 ("•  47 points total: A=11, B=7, C=9, D=6, E=6, F=3 (role-specific), G=5.", False),
 ("•  Yellow cells are for input. All other cells are formulas — do not overwrite them.", False),
 ("•  Correct answers and module references are shown on the Score Sheet and mirror the trainer Answer Key.", False),
]
r = 5
for text, bold in lines:
    c = ws.cell(row=r, column=2, value=text)
    c.font = Font(name=AR, bold=bold, size=11, color=(NAVY if bold else "000000"))
    r += 1
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 115

# ============ TAB 2: Score Sheet ============
ss = wb.create_sheet("Score Sheet")
ss.sheet_view.showGridLines = False
ss.merge_cells("A1:G1")
ss["A1"] = "CharmEd Minds™ Competency — Score Sheet"; ss["A1"].font = Font(name=AR, bold=True, size=15, color=NAVY)
ss.merge_cells("A2:G2")
ss["A2"] = "Mark each item 1 (correct) or 0 (incorrect) in the yellow column."; ss["A2"].font = Font(name=AR, italic=True, size=9, color="555555")

# input header block
def inp(cell_label, label_cell, input_cell):
    ss[label_cell] = cell_label; ss[label_cell].font = bfont
    ss[input_cell].fill = yfill; ss[input_cell].border = border; ss[input_cell].font = lfont
inp("Employee:", "A4", "B4"); inp("Role:", "D4", "E4")
inp("Date:", "A5", "B5"); inp("Grader:", "D5", "E5")

# table header
hrow = 7
headers = ["Item","Part","Question","Correct Answer","Module","Max","Awarded"]
for j, h in enumerate(headers, start=1):
    c = ss.cell(row=hrow, column=j, value=h)
    c.fill = hfill; c.font = hfont; c.border = border
    c.alignment = ctr if h in ("Item","Part","Module","Max","Awarded") else Alignment(horizontal="left", vertical="center")

dv = DataValidation(type="whole", operator="between", formula1=0, formula2=1, allow_blank=True,
                    errorTitle="Invalid", error="Enter 1 (correct) or 0 (incorrect).")
ss.add_data_validation(dv)

r0 = hrow + 1
for i, (item, part, q, ans, mod, mx) in enumerate(items):
    r = r0 + i
    vals = [item, part, q, ans, mod, mx, None]
    for j, v in enumerate(vals, start=1):
        c = ss.cell(row=r, column=j, value=v)
        c.border = border
        c.font = lfont
        if j in (3,4): c.alignment = wrap
        else: c.alignment = ctr
        if i % 2:
            if j <= 6: c.fill = zfill
    aw = ss.cell(row=r, column=7)
    aw.fill = yfill; dv.add(aw)
r1 = r0 + len(items) - 1

# ---- summary block ----
sr = r1 + 2
ss.cell(row=sr, column=1, value="SUMMARY").font = Font(name=AR, bold=True, size=12, color=NAVY)
sh = sr + 1
for j, h in enumerate(["Part","Section","Earned","Max","%"], start=1):
    c = ss.cell(row=sh, column=j, value=h); c.fill = hfill; c.font = hfont; c.border = border; c.alignment = ctr if h!="Section" else Alignment(horizontal="left")
awarded_rng = f"G{r0}:G{r1}"; part_rng = f"B{r0}:B{r1}"; max_rng = f"F{r0}:F{r1}"
row = sh + 1
part_rows = {}
for pt in parts:
    ss.cell(row=row, column=1, value=pt).alignment = ctr; ss.cell(row=row,column=1).border=border; ss.cell(row=row,column=1).font=bfont
    ss.cell(row=row, column=2, value=part_names[pt]).font = lfont; ss.cell(row=row,column=2).border=border
    ss.cell(row=row, column=3, value=f'=SUMIFS({awarded_rng},{part_rng},"{pt}")').border=border
    ss.cell(row=row, column=3).alignment=ctr; ss.cell(row=row,column=3).font=lfont
    ss.cell(row=row, column=4, value=f'=SUMIFS({max_rng},{part_rng},"{pt}")').border=border
    ss.cell(row=row, column=4).alignment=ctr; ss.cell(row=row,column=4).font=lfont
    ss.cell(row=row, column=5, value=f'=IFERROR(C{row}/D{row},0)').border=border
    ss.cell(row=row, column=5).number_format="0%"; ss.cell(row=row,column=5).alignment=ctr; ss.cell(row=row,column=5).font=lfont
    part_rows[pt]=row
    row += 1
# totals
trow = row
ss.cell(row=trow, column=2, value="TOTAL").font = Font(name=AR, bold=True, size=11)
ss.cell(row=trow, column=3, value=f"=SUM(C{sh+1}:C{trow-1})").font=bfont; ss.cell(row=trow,column=3).alignment=ctr; ss.cell(row=trow,column=3).border=border
ss.cell(row=trow, column=4, value=f"=SUM(D{sh+1}:D{trow-1})").font=bfont; ss.cell(row=trow,column=4).alignment=ctr; ss.cell(row=trow,column=4).border=border
ss.cell(row=trow, column=5, value=f"=IFERROR(C{trow}/D{trow},0)").font=bfont; ss.cell(row=trow,column=5).number_format="0%"; ss.cell(row=trow,column=5).alignment=ctr; ss.cell(row=trow,column=5).border=border
ss.cell(row=trow,column=1).border=border; ss.cell(row=trow,column=2).border=border

# result rows
br = trow + 2
ss.cell(row=br, column=2, value="Part B (Safety) — all correct?").font = bfont
ss.cell(row=br, column=3, value=f'=IF(C{part_rows["B"]}=D{part_rows["B"]},"YES","NO")').font=bfont
ss.cell(row=br, column=3).alignment=ctr; ss.cell(row=br,column=3).border=border; ss.cell(row=br,column=3).fill=zfill
rr = br + 1
ss.cell(row=rr, column=2, value="RESULT").font = Font(name=AR, bold=True, size=12, color=NAVY)
ss.cell(row=rr, column=3, value=f'=IF(AND(E{trow}>=0.8,C{part_rows["B"]}=D{part_rows["B"]}),"COMPETENT","NEEDS REVIEW")')
ss.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=5)
ss.cell(row=rr, column=3).font = Font(name=AR, bold=True, size=12); ss.cell(row=rr,column=3).alignment=ctr; ss.cell(row=rr,column=3).border=border
note = rr + 2
ss.cell(row=note, column=2, value="Pass = ≥80% overall AND Part B = 100%. If not, review Modules 3 & 8 and re-test.").font = Font(name=AR, italic=True, size=9, color="555555")

# conditional formatting for RESULT + Part B
res_cell = f"C{rr}"
ss.conditional_formatting.add(res_cell, CellIsRule(operator="equal", formula=['"COMPETENT"'], fill=PatternFill("solid",fgColor=GREEN), font=Font(name=AR,bold=True,color=GREENT)))
ss.conditional_formatting.add(res_cell, CellIsRule(operator="equal", formula=['"NEEDS REVIEW"'], fill=PatternFill("solid",fgColor=RED), font=Font(name=AR,bold=True,color=REDT)))
pb_cell = f"C{br}"
ss.conditional_formatting.add(pb_cell, CellIsRule(operator="equal", formula=['"NO"'], fill=PatternFill("solid",fgColor=RED), font=Font(name=AR,bold=True,color=REDT)))
ss.conditional_formatting.add(pb_cell, CellIsRule(operator="equal", formula=['"YES"'], fill=PatternFill("solid",fgColor=GREEN), font=Font(name=AR,bold=True,color=GREENT)))

widths = {"A":6,"B":7,"C":52,"D":58,"E":9,"F":6,"G":10}
for col,w in widths.items(): ss.column_dimensions[col].width = w
ss.freeze_panes = f"A{r0}"

# ============ TAB 3: Class Roster ============
cr = wb.create_sheet("Class Roster")
cr.sheet_view.showGridLines = False
cr.merge_cells("A1:N1")
cr["A1"] = "CharmEd Minds™ Competency — Class Roster"; cr["A1"].font = Font(name=AR, bold=True, size=15, color=NAVY)
cr.merge_cells("A2:N2")
cr["A2"] = "Enter each employee’s per-part points (yellow). Result = ≥80% overall AND Part B = 7/7."; cr["A2"].font = Font(name=AR, italic=True, size=9, color="555555")
rheaders = ["Employee","Role","Date","A /11","B /7","C /9","D /6","E /6","F /3","G /5","Total /47","%","Part B OK","Result"]
maxpts = {"D":11,"E":7,"F":9,"G":6,"H":6,"I":3,"J":5}  # columns D..J
hr = 4
for j,h in enumerate(rheaders, start=1):
    c = cr.cell(row=hr, column=j, value=h); c.fill=hfill; c.font=hfont; c.border=border
    c.alignment = ctr if j>=4 else Alignment(horizontal="left", vertical="center")
# example row + blanks
first = hr+1; nrows = 18
example = ["e.g., Jane Doe","Care Coordinator","2026-07-30",11,7,9,6,5,3,5]
for i in range(nrows):
    r = first + i
    for j in range(1, 15):
        c = cr.cell(row=r, column=j); c.border=border; c.font=lfont
        col = get_column_letter(j)
        if j<=3:
            c.alignment = Alignment(horizontal="left")
            if j<=3 and 4<=j: pass
        else:
            c.alignment = ctr
        if 1<=j<=10:  # input columns A..J
            c.fill = yfill
    # formulas
    cr.cell(row=r, column=11, value=f"=IF(COUNT(D{r}:J{r})=0,\"\",SUM(D{r}:J{r}))")  # Total
    cr.cell(row=r, column=12, value=f"=IF(K{r}=\"\",\"\",K{r}/47)"); cr.cell(row=r,column=12).number_format="0%"
    cr.cell(row=r, column=13, value=f'=IF(E{r}="","",IF(E{r}=7,"Yes","No"))')
    cr.cell(row=r, column=14, value=f'=IF(K{r}="","",IF(AND(L{r}>=0.8,E{r}=7),"Competent","Needs review"))')
    if i == 0:
        for j,v in enumerate(example, start=1):
            cc = cr.cell(row=r, column=j, value=v)
            cc.font = Font(name=AR, italic=True, size=10, color="777777")
# conditional formatting on Result column N
rng = f"N{first}:N{first+nrows-1}"
cr.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Competent"'], fill=PatternFill("solid",fgColor=GREEN), font=Font(name=AR,color=GREENT)))
cr.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Needs review"'], fill=PatternFill("solid",fgColor=RED), font=Font(name=AR,color=REDT)))
pbrng = f"M{first}:M{first+nrows-1}"
cr.conditional_formatting.add(pbrng, CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid",fgColor=RED), font=Font(name=AR,color=REDT)))
cwidths = {"A":18,"B":22,"C":12,"D":7,"E":7,"F":7,"G":7,"H":7,"I":7,"J":7,"K":10,"L":8,"M":11,"N":15}
for col,w in cwidths.items(): cr.column_dimensions[col].width = w
cr.freeze_panes = f"A{first}"

out = "CharmEd Minds Competency Scoring Grid.xlsx"
wb.save(out)
print("wrote", out)
